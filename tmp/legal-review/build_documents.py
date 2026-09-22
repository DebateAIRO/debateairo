from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path('/Users/vladmihaimiron/Documents/DebateAIRO')
WORK = BASE / 'tmp/legal-review'
OUT = BASE / 'output/legal-review'
OUT.mkdir(parents=True, exist_ok=True)

SOURCES = [
('OpenAI API data controls', 'https://developers.openai.com/api/docs/guides/your-data', 'Training, retention, ZDR and regional controls.'),
('OpenAI Services Agreement', 'https://openai.com/policies/services-agreement/', 'Contracting entity, commercial service and output terms.'),
('Anthropic Commercial Terms of Service', 'https://www.anthropic.com/legal/commercial-terms', 'Regional contracting entities and commercial terms.'),
('General Data Protection Regulation', 'https://eur-lex.europa.eu/eli/reg/2016/679/oj', 'Articles 3, 5, 6, 9, 13–14, 17, 22, 28 and 44–49; portal/indexed text.'),
('Romania Law 193 of 2000', 'https://legislatie.just.ro/Public/DetaliiDocumentAfis/77879', 'Unfair terms in consumer contracts.'),
('Romania OUG 141 of 2021', 'https://legislatie.just.ro/Public/DetaliiDocument/250054', 'Digital-service conformity, remedies and modifications; indexed text.'),
('Romania OUG 34 of 2014', 'https://legislatie.just.ro/Public/DetaliiDocument/158913', 'Consumer information, withdrawal and Article 11¹ online function.'),
('Digital Services Act', 'https://eur-lex.europa.eu/eli/reg/2022/2065/oj', 'Hosting and platform framework; Articles 3, 6, 11–21; portal/indexed text.'),
('ANCOM guidance for intermediary services', 'https://www.ancom.ro/servicii-digitale/furnizori-servicii-intermediare/', 'Romanian DSA supervision and differentiated obligations.'),
('European Commission Article 50 transparency FAQ', 'https://digital-strategy.ec.europa.eu/en/faqs/transparency-obligations-under-article-50-ai-act', 'AI transparency duties and 2026 transition dates.'),
('Gemini API Additional Terms of Service', 'https://ai.google.dev/gemini-api/terms', 'Paid/unpaid data use, regional conditions and Search grounding restrictions.'),
('EDPB international transfer guidance', 'https://www.edpb.europa.eu/sme/be-compliant/international-data-transfers_en', 'Adequacy, safeguards and transfer-risk analysis.'),
('Romania Law 365 of 2002', 'https://legislatie.just.ro/Public/DetaliiDocument/37075', 'Trader information and electronic contracting.'),
('German Civil Code section 312k', 'https://www.gesetze-im-internet.de/bgb/__312k.html', 'Online cancellation of qualifying consumer contracts.'),
('German Civil Code section 309', 'https://www.gesetze-im-internet.de/bgb/__309.html', 'Standard-term restrictions, including renewal terms in paragraph 9.'),
('France Service Public subscription cancellation', 'https://www.service-public.gouv.fr/particuliers/vosdroits/F33991', 'Renewal and online cancellation of service contracts.'),
('ICO UK IDTA and Addendum guidance', 'https://ico.org.uk/for-organisations/uk-gdpr-guidance-and-resources/international-transfers/appropriate-safeguards/what-are-standard-data-protection-clauses-the-uk-idta-and-the-addendum/', 'UK transfer instruments and their use.'),
('Ofcom AI chatbots and online regulation', 'https://www.ofcom.org.uk/online-safety/illegal-and-harmful-content/ai-chatbots-and-online-regulation-what-you-need-to-know', 'User-shared AI content and Online Safety Act scope.'),
('UK government subscription regime response', 'https://www.gov.uk/government/consultations/consultation-on-the-implementation-of-the-new-subscription-contracts-regime/outcome/government-response-to-consultation-on-the-implementation-of-the-new-subscription-contracts-regime-web-accessible-version', 'Anticipated spring 2027 commencement and proposed implementation.'),
('California Privacy Protection Agency threshold adjustment', 'https://cppa.ca.gov/announcements/2024/20241217.html', 'Adjusted annual gross-revenue threshold of $26,625,000.'),
('California Attorney General CCPA guidance', 'https://www.oag.ca.gov/privacy/ccpa', 'Coverage, consumer rights, notices and business obligations.'),
('FTC ROSCA guidance', 'https://www.ftc.gov/business-guidance/blog/2018/07/time-rosca-recap-ftc-says-risk-free-trial-was-risky-not-free', 'Disclosure, consent and cancellation for online negative options.'),
('California Business and Professions Code section 17602', 'https://leginfo.legislature.ca.gov/faces/codes_displaySection.xhtml?lawCode=BPC&sectionNum=17602.', 'Automatic-renewal requirements and July 2025 contract applicability.'),
('Canada OPC cross-border processing guidance', 'https://www.priv.gc.ca/en/privacy-topics/airports-and-borders/gl_dab_090127/', 'PIPEDA accountability, contracts and transparency for outsourcing.'),
('Quebec private-sector privacy law', 'https://www.legisquebec.gouv.qc.ca/fr/document/lc/P-39.1', 'Section 17 assessments and agreements for disclosure outside Quebec.'),
('Quebec OQLF adhesion-contract guidance', 'https://www.oqlf.gouv.qc.ca/francisation/entreprises/contrats-adhesion.html', 'French-language provision and exceptions.'),
('Australia ACCC consumer rights and guarantees', 'https://www.accc.gov.au/consumers/buying-products-and-services/consumer-rights-and-guarantees', 'Statutory service guarantees and overseas suppliers.'),
('Australia OAIC APP 8 guidance', 'https://www.oaic.gov.au/privacy/australian-privacy-principles/australian-privacy-principles-guidelines/chapter-8-app-8-cross-border-disclosure-of-personal-information', 'Cross-border disclosure safeguards and accountability.'),
('Australia OAIC small-business guidance', 'https://www.oaic.gov.au/privacy/privacy-guidance-for-organisations-and-government-agencies/organisations/small-business', 'Privacy Act coverage and small-business exceptions.'),
('Brazil Consumer Defence Code', 'https://www.planalto.gov.br/ccivil_03/leis/l8078compilado.htm', 'Article 49 distance-contract withdrawal.'),
('Brazil ANPD international transfer guidance', 'https://www.gov.br/anpd/pt-br/assuntos/assuntos-internacionais/transferencia-internacional-de-dados/international-affairs', 'Resolution 19/2024, EU adequacy and status of equivalent clauses.'),
('EU adequacy decision for Brazil 2026', 'https://eur-lex.europa.eu/legal-content/EN/ALL/?uri=OJ%3AL_202600179', 'Commission Implementing Decision (EU) 2026/179.'),
('OpenAI Data Processing Addendum', 'https://openai.com/policies/data-processing-addendum/', 'January 2026 DPA and regional parties.'),
('Anthropic Data Processing Addendum', 'https://www.anthropic.com/legal/data-processing-addendum', 'Processor arrangements and transfer provisions.'),
('Anthropic commercial retention policy', 'https://privacy.claude.com/en/articles/7996866-how-long-do-you-store-my-organization-s-data', 'Default API retention and exceptions.'),
('Claude API and data retention', 'https://platform.claude.com/docs/en/manage-claude/api-and-data-retention', 'ZDR, covered-model requirements and feature exceptions.'),
('Google contracting entity table', 'https://cloud.google.com/terms/google-entity', 'Contracting entities by billing location and service.'),
('China Personal Information Protection Law', 'https://www.cac.gov.cn/2021-08/20/c_1631050028355286.htm', 'Articles 3, 13, 17, 28–29, 38–40, 53 and 55; CAC text.'),
('China cross-border data flow provisions', 'https://www.cac.gov.cn/2024-03/22/c_1712776611775634.htm', '2024 export-route exemptions, thresholds and continuing privacy duties.'),
('China Generative AI Interim Measures', 'https://www.cac.gov.cn/2023-07/13/c_1690898327029107.htm', 'Public service scope; content, privacy, complaints and Article 17 filing.'),
('China AI generated content labeling measures', 'https://www.cac.gov.cn/2025-03/14/c_1743654684782215.htm', 'Visible and metadata labels, exports and Article 8 user agreements.'),
('China consumer protection implementation regulation', 'https://app.www.gov.cn/govdata/gov/202403/19/513111/article.html', '2024 regulation; renewal disclosure and consumer contract protection.'),
('Japan Act on the Protection of Personal Information', 'https://www.japaneselawtranslation.go.jp/en/laws/view/4241/en', 'PPC-linked official translation; purpose, transfers and overseas scope.'),
('Japan PPC Supplementary Rules', 'https://www.ppc.go.jp/files/pdf/Supplementary_Rules_en.pdf', 'EU/UK-origin data under adequacy and onward-transfer safeguards.'),
('Japan CAA mail-order transaction guidance', 'https://www.no-trouble.caa.go.jp/what/mailorder/', 'Seller and service disclosures; online order confirmation.'),
('European Commission adequacy decisions', 'https://commission.europa.eu/law/law-topic/data-protection/international-dimension-data-protection/adequacy-decisions_en', 'Current adequacy list, including Japan and Korea; indexed text.'),
('Singapore PDPC data protection obligations', 'https://www.pdpc.gov.sg/overview-of-pdpa/the-legislation/personal-data-protection-act/data-protection-obligations', 'Accountability, retention and comparable overseas protection; indexed text.'),
('Singapore PDPC GenAI personal-data guidelines', 'https://files.app.optical.gov.sg/pdpc/production/assets/143cb9d4-532e-4cca-9a77-bcc0415ca294.pdf', 'Final guidelines issued 20 July 2026; training notices and stakeholder roles.'),
('Singapore CCS subscription-trap enforcement', 'https://www.ccs.gov.sg/media-and-events/newsroom/announcements-and-media-releases/e-commerce-retailer-fashion-interactive-ordered-to-cease-unfair-trade-practices-and-stop-using--subscription-traps--481/', 'Unfair recurring-subscription enrolment and consumer protection.'),
('Singapore cancellation of contracts regulations', 'https://sso.agc.gov.sg/SL/CPFTA2003-RG1?DocDate=20241217&ValidDate=20241218', 'Specified regulated contract types; indexed official text.'),
('South Korea Personal Information Protection Act', 'https://elaw.klri.re.kr/eng_service/lawViewContent.do?hseq=71740', 'Articles 23, 26, 28-8, 28-11 and 30; official English translation.'),
('Korea PIPC foreign business guidance', 'https://pipc.go.kr/eng/user/ltn/new/noticeDetail.do?bbsId=BBSMSTR_000000000001&nttId=2488', 'Overseas scope, entrusted processing, notices and domestic-agent scoping.'),
('Korea MSIT AI transparency guidance', 'https://www.msit.go.kr/eng/bbs/view.do?bbsSeqNo=42&nttSeqNo=1215', 'January 2026 guidance on advance notice and in-service/exported labels.'),
('Korea MSIT AI Basic Act implementation', 'https://www.msit.go.kr/eng/bbs/view.do?bbsSeqNo=42&nttSeqNo=1214', '22 January 2026 commencement and announced enforcement grace period.'),
('Korea consumer protection in electronic commerce', 'https://elaw.klri.re.kr/eng_mobile/subjectViewer.do?hseq=69592&key=08&pCode=126&pName=Online+Shopping&type=subject', 'Articles 5, 13 and 17–18; cancellation and recurring-payment changes.'),
('OpenAI API supported territories', 'https://help.openai.com/en/articles/5347006-openai-api-supported-countries-and-territories', 'Availability and restrictions on offering access outside listed territories.'),
('Anthropic supported countries and regions', 'https://www.anthropic.com/supported-countries', 'Commercial API territory list and ownership-related conditions.'),
('Gemini API available regions', 'https://ai.google.dev/gemini-api/docs/available-regions', 'Direct Gemini API and AI Studio territory list.'),
('Japan CAA consumer transaction guide in English', 'https://www.no-trouble.caa.go.jp/foreignlanguage/english/', 'Mail-order exclusion from general cooling-off; other remedies.'),
('Japan Cabinet Office AI guidance', 'https://www8.cao.go.jp/cstp/ai/ai_guideline/ai_guideline.html', 'December 2025 national guidance and March 2026 business guideline links.'),
]

def norm(text):
    return text.replace('–', '-').replace('—', '-').replace('\u2011', '-')

def hyperlink(p, label, url):
    el = OxmlElement('w:hyperlink')
    el.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '174C78'); pr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); pr.append(u)
    r.append(pr)
    t = OxmlElement('w:t'); t.text = norm(label); r.append(t)
    el.append(r); p._p.append(el)

def inline(p, text):
    parts = re.split(r'(\*\*.*?\*\*|\[S\d{2}\]|https://[^\s]+)', text)
    for chunk in parts:
        if not chunk:
            continue
        if chunk.startswith('**') and chunk.endswith('**'):
            p.add_run(norm(chunk[2:-2])).bold = True
        elif re.fullmatch(r'\[S\d{2}\]', chunk):
            idx=int(chunk[2:4])-1
            hyperlink(p, chunk, SOURCES[idx][1])
        elif chunk.startswith('https://'):
            url=chunk.rstrip('.,;')
            hyperlink(p,url,url)
            if chunk[len(url):]: p.add_run(chunk[len(url):])
        else:
            p.add_run(norm(chunk))

def field(p, instr):
    r=p.add_run()
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), instr)
    r._r.addnext(fld)

def setup(title, short):
    d=Document()
    sec=d.sections[0]
    sec.page_width=Inches(8.2677); sec.page_height=Inches(11.6929)
    sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.67)
    sec.left_margin=Inches(.77); sec.right_margin=Inches(.77)
    sec.header_distance=Inches(.3); sec.footer_distance=Inches(.3)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=d.styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        rf=st._element.get_or_add_rPr().rFonts
        for key in list(rf.attrib):
            if 'Theme' in key or 'theme' in key: del rf.attrib[key]
        for border in list(st._element.xpath('./w:pPr/w:pBdr')):border.getparent().remove(border)
        for col in st._element.xpath('./w:rPr/w:color'):
            for key in list(col.attrib):
                if 'theme' in key.lower():del col.attrib[key]
    n=d.styles['Normal']; n.font.size=Pt(11)
    n.paragraph_format.line_spacing=1.08
    n.paragraph_format.space_after=Pt(6)
    n.paragraph_format.widow_control=True
    n.paragraph_format.keep_together=True
    for name,size,before,after in [('Title',23,0,9),('Heading 1',14,13,7),('Heading 2',12,9,5)]:
        st=d.styles[name]; st.font.size=Pt(size); st.font.bold=name!='Title'
        st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after)
        st.paragraph_format.keep_with_next=True
    h=sec.header.paragraphs[0]; h.text=short+' | Review draft'
    h.style=d.styles['Normal']
    for r in h.runs:r.font.size=Pt(9)
    h.paragraph_format.space_after=Pt(0)
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run('21 September 2026  |  Page ');field(f,'PAGE');f.add_run(' of ');field(f,'NUMPAGES')
    for r in f.runs:r.font.size=Pt(9)
    d.core_properties.title=title
    d.core_properties.subject='Draft legal review for DebateAIRO'
    d.core_properties.author=''
    d.core_properties.keywords='DebateAI, draft, terms, country research'
    return d

def table(d, lines):
    rows=[[c.strip() for c in l.strip().strip('|').split('|')] for l in lines]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',x) for x in r)]
    t=d.add_table(rows=0, cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[1.75,4.98] if len(rows[0])==2 else [6.73/len(rows[0])]*len(rows[0])
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    pr.append(borders)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,(c,tx) in enumerate(zip(cells,row)):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'95');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E9EEF3' if i==0 else 'FFFFFF');tcpr.append(shading)
            p=c.paragraphs[0];inline(p,tx);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.04
            for r in p.runs:r.font.size=Pt(10);r.bold=i==0
        trpr=t.rows[-1]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(2)

def build(src, title, short, path):
    text=src.read_text()
    refs=re.findall(r'\[S(\d+)\]',text)
    assert all(1<=int(r)<=len(SOURCES) for r in refs)
    d=setup(title,short)
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            group=[]
            while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i]);i+=1
            table(d,group);continue
        if line=='{{SOURCES}}':
            for n,(label,url,desc) in enumerate(SOURCES,1):
                p=d.add_paragraph();p.paragraph_format.space_after=Pt(6)
                p.paragraph_format.keep_together=True
                p.add_run(f'[S{n:02}] ')
                hyperlink(p,label,url)
                p.add_run('. '+norm(desc))
                for r in p._p.xpath('.//w:r'):
                    pr=r.find(qn('w:rPr'))
                    if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
                    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');pr.append(sz)
            i+=1;continue
        if line.startswith('# '):
            d.add_paragraph(norm(line[2:]),style='Title')
        elif line.startswith('## '):
            if line=='## Sources checked':
                sec=d.add_section(WD_SECTION.NEW_PAGE)
                cols=sec._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'360')
            h=d.add_paragraph(norm(line[3:]),style='Heading 1')
        elif line.startswith('### '):d.add_paragraph(norm(line[4:]),style='Heading 2')
        else:
            p=d.add_paragraph();inline(p,line)
        i+=1
    d.save(path)
    print(f'{path}: {len(text.split())} words, {len(d.paragraphs)} paragraphs, {len(d.tables)} tables')

build(WORK/'revised-terms.md','DebateAI Terms of Service','DebateAI Terms',OUT/'DebateAI_Revised_Terms_Expanded.docx')
build(WORK/'review.md','DebateAI Terms Review and Country Requirements','DebateAI Legal Review',WORK/'DebateAI_Review_and_Country_Research_Expanded.docx')
